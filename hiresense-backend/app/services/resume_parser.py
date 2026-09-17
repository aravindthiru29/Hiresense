import os
from typing import Tuple, Optional
from app.utils.helpers import clean_text


class ResumeParser:
    """Extracts raw text from PDF and DOCX files."""

    @staticmethod
    def extract_text_from_file(file_path: str) -> Tuple[Optional[str], Optional[str]]:
        if not os.path.exists(file_path):
            return None, "File not found on server"

        ext = file_path.rsplit('.', 1)[-1].lower()

        try:
            if ext == 'pdf':
                return ResumeParser._extract_pdf(file_path), None
            elif ext in ('docx', 'doc'):
                return ResumeParser._extract_docx(file_path), None
            else:
                return None, f"Unsupported file format: .{ext}. Supported formats: PDF, DOCX"
        except Exception as e:
            return None, f"Error reading document: {str(e)}"

    @staticmethod
    def _extract_pdf(file_path: str) -> str:
        text_content = []
        try:
            import fitz  # PyMuPDF
            doc = fitz.open(file_path)
            for page in doc:
                text_content.append(page.get_text())
            doc.close()
        except ImportError:
            # Fallback to pdfminer if fitz is somehow unavailable
            from pdfminer.high_level import extract_text
            return clean_text(extract_text(file_path))

        full_text = "\n".join(text_content)
        return clean_text(full_text)

    @staticmethod
    def _extract_docx(file_path: str) -> str:
        import docx
        doc = docx.Document(file_path)
        text_content = [paragraph.text for paragraph in doc.paragraphs if paragraph.text]
        for table in doc.tables:
            for row in table.rows:
                row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                if row_text:
                    text_content.append(row_text)
        return clean_text("\n".join(text_content))
